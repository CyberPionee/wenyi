"""Offline source parsing must not initialize translation or call models."""

from unittest.mock import patch

import pytest
from typer.testing import CliRunner
from wenyi_cli.cli import app
from wenyi_core.config import Config
from wenyi_core.ingest.fb2_reader import peek_fb2_title, read_fb2
from wenyi_core.ingest.models import Chapter, Document, Segment
from wenyi_core.llm.providers.fake import FakeClient
from wenyi_core.pipeline import input_preparation
from wenyi_core.pipeline.orchestrator import Orchestrator
from wenyi_core.storage.file import FileStorage

from tests.fake_llm import routing_handler


def test_cli_parse_without_credentials_or_orchestrator(tmp_path, monkeypatch):
    source = tmp_path / "book.txt"
    source.write_text("# First chapter\n\nA short story.\n", encoding="utf-8")
    config = Config.from_dict(
        {
            "language": {"source": "auto", "target": "zh"},
            "paths": {"state_dir": str(tmp_path / "state")},
        }
    )
    monkeypatch.setattr(
        "wenyi_cli.commands.context.CommandContext.load_config", lambda self: config
    )
    with (
        patch("wenyi_core.config.Config.create_default_file"),
        patch("wenyi_core.pipeline.orchestrator.Orchestrator", side_effect=AssertionError),
        patch(
            "wenyi_cli.commands.context.CommandContext.validate_api_configuration",
            side_effect=AssertionError,
        ),
    ):
        result = CliRunner().invoke(app, ["parse", str(source)])
    assert result.exit_code == 0, result.output
    assert "First chapter" in result.output
    assert "State directory:" in result.output
    artifacts = list((tmp_path / "state").rglob("parsed_document.json"))
    assert len(artifacts) == 1
    store = FileStorage(str(artifacts[0].parent))
    assert not store.exists()
    assert store.load_analysis() is None
    parsed = store.read_artifact("parsed_document.json")
    assert parsed["source_sha256"]
    assert parsed["ingest_config"]["source_lang"] == "auto"
    assert any(
        segment["source"] == "A short story."
        for chapter in parsed["document"]["chapters"]
        for segment in chapter["segments"]
    )
    assert not list((tmp_path / "state").rglob(".initializing.json"))


def inputs(tmp_path):
    source = tmp_path / "book.txt"
    source.write_text("# First chapter\n\nA short story.\n", encoding="utf-8")
    config = Config.from_dict(
        {
            "language": {"source": "en", "target": "zh"},
            "llm": {"preset": "fake"},
            "pipeline": {"book_understanding": False},
            "paths": {"state_dir": str(tmp_path / "state")},
        }
    )
    return source, config, FileStorage(str(tmp_path / "run"))


@pytest.mark.parametrize("extension", [".txt", ".docx", ".fb2"])
@pytest.mark.parametrize("source_lang", ["en", "auto"])
def test_local_parse_then_prepare_reuses_cache(tmp_path, extension, source_lang):
    source, config, _ = inputs(tmp_path)
    config.source_lang = source_lang
    source = source.with_suffix(extension)
    if extension == ".docx":
        from docx import Document as DocxDocument

        document = DocxDocument()
        document.add_heading("First chapter", 1)
        document.add_paragraph("A short story.")
        document.save(source)
    elif extension == ".fb2":
        source.write_text(
            '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">'
            "<description><title-info><book-title>Metadata title</book-title>"
            "</title-info></description><body><section><p>A short story.</p>"
            "</section></body></FictionBook>",
            encoding="utf-8",
        )
    with patch.object(
        input_preparation, "load_document", wraps=input_preparation.load_document
    ) as load:
        store, digest = input_preparation.locate_input_storage(str(source), config, FileStorage)
        with store.lock():
            parsed = input_preparation.parse_document(
                store, str(source), config, expected_sha256=digest
            )
        load.assert_called_once()
    with patch.object(input_preparation, "load_document", side_effect=AssertionError("reparsed")):
        located, digest = input_preparation.locate_input_storage(str(source), config, FileStorage)
        with located.lock():
            assert (
                input_preparation.parse_document(
                    located, str(source), config, expected_sha256=digest
                )
                == parsed
            )
        prepared = Orchestrator(config, client=FakeClient(handler=routing_handler)).prepare(
            str(source)
        )
    assert prepared.run_dir == store.run_dir
    assert prepared.load_manifest()["title"] == parsed.title
    assert prepared.load_manifest()["source_lang"] == ("ja" if source_lang == "auto" else "en")
    source.write_bytes(source.read_bytes().replace(b"A short story.", b"A changed story."))
    # DOCX is a ZIP archive; appending bytes changes identity without corrupting its title.
    if extension == ".docx":
        source.write_bytes(source.read_bytes() + b"changed")
    with patch.object(input_preparation, "load_document", side_effect=AssertionError("reparsed")):
        with pytest.raises(ValueError, match="does not match"):
            Orchestrator(config, client=FakeClient(handler=routing_handler)).prepare(str(source))


@pytest.mark.parametrize("namespace", ["", "http://www.gribuser.ru/xml/fictionbook/2.0", "urn:2.1"])
@pytest.mark.parametrize("encoding", ["utf-8", "windows-1251", "unknown-encoding"])
@pytest.mark.parametrize(
    ("metadata", "expected"),
    [
        ("<title-info><book-title> Книга </book-title></title-info>", "Книга"),
        ("<title-info><book-title /></title-info>", "book"),
        ("<title-info><book-title>  </book-title></title-info>", ""),
        (
            "<title-info><book-title>First<emphasis>ignored</emphasis></book-title></title-info>"
            "<title-info><book-title>Last</book-title></title-info>",
            "Last",
        ),
    ],
)
def test_fb2_metadata_title_matches_parser(tmp_path, namespace, encoding, metadata, expected):
    source = tmp_path / "book.fb2"
    xml = (
        f'<?xml version="1.0" encoding="{encoding}"?>'
        f'<FictionBook xmlns="{namespace}"><description>{metadata}</description>'
        "<body><section><p>A story.</p></section></body></FictionBook>"
    )
    source.write_bytes(xml.encode("utf-8" if encoding == "unknown-encoding" else encoding))
    assert peek_fb2_title(str(source)) == read_fb2(str(source), "en", "zh").title == expected


@pytest.mark.parametrize(
    "extension", [".txt", ".text", ".md", ".markdown", ".html", ".htm", ".xhtml"]
)
def test_filename_location_matches_parser(tmp_path, extension):
    source, config, _ = inputs(tmp_path)
    source = source.with_suffix(extension.upper())
    source.write_text("<html><title>Not the book title</title><body>A story.</body></html>")
    parsed = input_preparation.load_document(str(source), "en", "zh")
    with patch.object(input_preparation, "load_document", side_effect=AssertionError("parsed")):
        store, _ = input_preparation.locate_input_storage(str(source), config, FileStorage)
    assert parsed.title == source.stem
    assert store.run_dir == input_preparation.translation_run_dir(
        config.state_dir, parsed.title, config.target_lang
    )


def test_unknown_format_location_rejected(tmp_path):
    source, config, _ = inputs(tmp_path)
    unknown = source.with_suffix(".unknown")
    unknown.write_text("A story.")
    with pytest.raises(ValueError, match="Unsupported format: .unknown"):
        input_preparation.locate_input_storage(str(unknown), config, FileStorage)
    assert not (tmp_path / "state").exists()


def test_source_changed_during_metadata_lookup_rejected(tmp_path):
    source, config, _ = inputs(tmp_path)
    source = source.with_suffix(".fb2")
    source.write_text("Original source")

    def changing_title(path):
        source.write_text("Changed source")
        return "book"

    with patch.object(input_preparation, "peek_fb2_title", side_effect=changing_title):
        with pytest.raises(ValueError, match="Source changed"):
            input_preparation.locate_input_storage(str(source), config, FileStorage)
    assert not (tmp_path / "state").exists()


def test_local_title_location_does_not_parse_twice(tmp_path):
    source, config, _ = inputs(tmp_path)
    with patch.object(
        input_preparation, "load_document", wraps=input_preparation.load_document
    ) as load:
        store, digest = input_preparation.locate_input_storage(str(source), config, FileStorage)
        load.assert_not_called()
        with store.lock():
            parsed = input_preparation.parse_document(
                store, str(source), config, expected_sha256=digest
            )
    load.assert_called_once()
    assert parsed.title == source.stem
    assert store.run_dir == input_preparation.translation_run_dir(
        config.state_dir, parsed.title, config.target_lang
    )


def test_parse_cache_and_settings_invalidation(tmp_path):
    source, config, store = inputs(tmp_path)
    with (
        store.lock(),
        patch.object(
            input_preparation, "load_document", wraps=input_preparation.load_document
        ) as load,
    ):
        first = input_preparation.parse_document(store, str(source), config)
        assert input_preparation.parse_document(store, str(source), config) == first
        assert load.call_count == 1
        config.segment.max_tokens_per_segment += 1
        input_preparation.parse_document(store, str(source), config)
        assert load.call_count == 2
        source.write_text("# First chapter\n\nChanged story.", encoding="utf-8")
        changed = input_preparation.parse_document(store, str(source), config)
        assert changed != first
        assert load.call_count == 3
    assert not store.exists()


def test_source_changed_during_parse_does_not_publish(tmp_path, monkeypatch):
    source, config, store = inputs(tmp_path)
    load = input_preparation.load_document

    def changing_load(*args, **kwargs):
        document = load(*args, **kwargs)
        source.write_text("Changed content", encoding="utf-8")
        return document

    monkeypatch.setattr(input_preparation, "load_document", changing_load)
    with store.lock(), pytest.raises(ValueError, match="Source changed"):
        input_preparation.parse_document(store, str(source), config)
    assert store.read_artifact("parsed_document.json") is None
    assert not store.exists()


@pytest.mark.parametrize("pdf", [False, True])
def test_parse_then_prepare_reuses_artifact_and_preserves_initialized_state(tmp_path, pdf):
    source, config, store = inputs(tmp_path)
    if pdf:
        source = tmp_path / "book.pdf"
        source.write_bytes(b"%PDF-1.4 offline mock")
        document = Document(
            title="book",
            fmt="pdf",
            source_path=str(source),
            source_lang="en",
            target_lang="zh",
            chapters=[
                Chapter(index=0, title="First", segments=[Segment(index=0, source="A story.")])
            ],
        )
        with patch.object(input_preparation, "load_document", return_value=document) as load:
            store, digest = input_preparation.locate_input_storage(str(source), config, FileStorage)
            with store.lock():
                input_preparation.parse_document(store, str(source), config, expected_sha256=digest)
                input_preparation.parse_document(store, str(source), config)
            load.assert_called_once()
            assert load.call_args.kwargs["cache_dir"] == store.source_dir
            assert load.call_args.kwargs["source_hash"] == digest
            assert load.call_args.kwargs["pdf_backend"] == config.pipeline.pdf_backend
    else:
        with store.lock():
            input_preparation.parse_document(store, str(source), config)

    client = FakeClient(handler=routing_handler)
    with patch.object(input_preparation, "load_document", side_effect=AssertionError("reparsed")):
        Orchestrator(config, client=client, storage=None if pdf else store).prepare(str(source))
    assert store.exists()
    assert any(call["operation"] == "analysis.style" for call in client.calls)
    chapter = store.load_chapter(0)
    chapter.segments[-1].target = "Saved human translation"
    store.save_chapter(chapter)
    manifest = store.load_manifest()
    analysis = store.load_analysis()
    usage = store.load_usage()
    events = store.read_artifact_records("events.jsonl")
    with store.lock():
        parsed = input_preparation.parse_document(store, str(source), config)
    assert all(
        segment.target is None for chapter in parsed.chapters for segment in chapter.segments
    )
    assert store.load_manifest() == manifest
    assert store.load_chapter(0) == chapter
    assert store.load_analysis() == analysis
    assert store.load_usage() == usage
    assert store.read_artifact_records("events.jsonl") == events
    source.write_bytes(b"Different content with the same name")
    with store.lock(), pytest.raises(ValueError, match="does not match"):
        input_preparation.parse_document(store, str(source), config)
    assert store.load_manifest() == manifest
    assert store.load_chapter(0) == chapter


def test_cli_parse_rejects_subtitles_without_creating_state(tmp_path, monkeypatch):
    _, config, _ = inputs(tmp_path)
    source = tmp_path / "captions.srt"
    source.write_text("1\n00:00:00,000 --> 00:00:01,000\nHello\n", encoding="utf-8")
    monkeypatch.setattr(
        "wenyi_cli.commands.context.CommandContext.load_config", lambda self: config
    )
    with patch("wenyi_core.config.Config.create_default_file"):
        result = CliRunner().invoke(app, ["parse", str(source)])
    assert result.exit_code == 1
    assert "wenyi translate" in result.output
    assert not (tmp_path / "state").exists()
